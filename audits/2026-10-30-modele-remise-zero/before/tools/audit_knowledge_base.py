"""Audit read-only: active knowledge, provenance, and saved Excel versus YAML."""
import argparse
from collections import defaultdict
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
from urllib.parse import unquote

from validate_stacks import ROOT, read
from project_investment import prepare, project

WORKBOOK = ROOT/'outputs/01a15da9-73b2-7760-8111-3c718f6245f3/FLOW_Investissements_10_ans_Integration_Migration.xlsx'


def active_files():
    files = [ROOT/'AGENTS.md', ROOT/'README.md']
    for folder in ('source', 'sources', 'config'):
        files += [p for p in (ROOT/folder).rglob('*') if p.is_file() and p.suffix in ('.yaml', '.yml', '.md')]
    return sorted(files)


def inventory(files):
    groups=defaultdict(list)
    errors=[]; links=0
    for path in files:
        rel=path.relative_to(ROOT).as_posix()
        raw=path.read_bytes(); text=raw.decode('utf-8')
        groups[hashlib.sha256(raw).hexdigest()].append(rel)
        if re.search(r'Ã[©¨ª]|â€[™œ“]',text): errors.append(f'Encodage suspect : {rel}')
        if path.suffix in ('.yaml','.yml'): read(path)  # rejects duplicate YAML keys
        if path.suffix=='.md':
            for target in re.findall(r'!?\[[^\]]*\]\((<[^>]+>|[^)]+)\)',text):
                target=unquote(target.strip('<>')).split('#')[0]
                if not target or re.match(r'[a-zA-Z]+:',target): continue
                links+=1
                if not (path.parent/target).is_file(): errors.append(f'Lien absent : {rel} -> {target}')
    duplicates=[g for g in groups.values() if len(g)>1]
    if duplicates: errors.append('Fichiers actifs identiques : '+str(duplicates))
    for path in ROOT.glob('[0-9][0-9]-*.md'): errors.append('Ancienne note encore à la racine : '+path.name)
    return dict(files=len(files),local_links=links,exact_duplicates=duplicates),errors


def check_glossary_references(model):
    errors=[]
    for key,term in model['notions'].items():
        file,_,anchor=term['source_ref'].partition('#')
        path=ROOT/file
        if not path.is_file():
            errors.append(f'Notion {key} : source absente {file}')
            continue
        if anchor and path.suffix in ('.yaml','.yml'):
            value=read(path)
            try:
                for part in anchor.split('/'): value=value[int(part)] if isinstance(value,list) else value[part]
            except (KeyError,IndexError,ValueError,TypeError): errors.append(f'Notion {key} : ancre absente {term["source_ref"]}')
    return errors


def bindings(data):
    """Every numeric literal in Hypothèses and Sources has an explicit model binding."""
    out={}
    def bind(sheet,cell,value,ref): out[(sheet,cell)]=dict(value=value,reference=ref)
    def h(cell,value,ref): bind('Hypothèses',cell,value,ref)
    def s(cell,value,ref): bind('Sources',cell,value,ref)
    for i,o in enumerate(data['options']):
        c='CDEFG'[i]; base=o['base_annuelle']; ref=f"source/chiffrage/chiffrage-{o['stack']}.yaml"
        for row,value in {7:o['build']/1e6,8:o['integration_source']/1e6,9:base['services_mco']/1e6,
                          10:base['hebergement']/1e6,11:base['ipaas']/1e6,12:base['licences_complementaires']/1e6,
                          13:len(o['prix_abonnements']),19:None if o['ventilation_services']['core'] is None else o['ventilation_services']['core']/1e6,
                          20:o['build_mois'],21:o['integration_mois'] if i>=3 else None}.items(): s(f'{c}{row}',value,ref)
        for k in range(5): s(f'{c}{14+k}',o['prix_abonnements'][k]/1e6 if k<len(o['prix_abonnements']) else None,ref)
        for row,value in {8:o['build_mois'],9:o['build']/1e3,18:(o['ventilation_services']['core'] if i<3 else base['services_mco'])/1e3,
                          20:base['hebergement']/1e3,21:base['ipaas']/1e3,22:base['licences_complementaires']/1e3}.items(): h(f'{c}{row}',value,ref)
        if i<3: h(f'{c}19',o['ventilation_services']['integration']/1e3,'source/plan-investissement.yaml#presentation_achats')
        else:
            h(f'{c}13',o['integration_mois'],'source/plan-investissement.yaml#integration_buy_mois')
            h(f'{c}14',o['integration']/1e3,ref)
    h('C15',data['bonus_continuite_make']['taux_pct']/100,'source/plan-investissement.yaml#integration_make')
    for r,key in [(28,'tjm_ingenierie'),(29,'licences_logiciels'),(30,'cloud')]:
        h(f'C{r}',data['taux'][key],'source/inflation.yaml#categories')
        s(f'C{r-5}',data['taux'][key],'source/inflation.yaml#categories')
    for r,lot in zip((55,56),data['sensibilite_finance_sap']['lots']):
        for c,key in [('C','jh_externes'),('D','jh_finance_fpa'),('E','prix_lot_eur_ht')]: h(f'{c}{r}',lot[key]/(1000 if c=='E' else 1),'source/plan-investissement.yaml#sensibilite_finance_sap')
    a=data['sensibilite_finance_sap']; p=data['comparaison_prix_integration']
    h('G59',a['complements_sap_pilotage_preparation_retenus_dans_ancre_make_eur_ht']/1e3,'source/chiffrage/chiffrage-sap.yaml')
    h('G65',p['jh_base_sap_proxy'],'source/plan-investissement.yaml#sensibilite_finance_sap')
    roles=['integration_metier','integration_it_donnees_qualification','developpement_java_react','pilotage']
    for r,role in zip(range(67,71),roles):
        h(f'C{r}',p['profils_base'][role]['part_pct']/100,'source/plan-investissement.yaml#comparaison_prix_integration_miroir_sap')
        h(f'E{r}',p['profils_base'][role]['tjm_eur_ht'],'source/plan-investissement.yaml#comparaison_prix_integration_miroir_sap')
    comp=defaultdict(float)
    for row in p['complements_jh_par_nature'].values(): comp[row['profil']]+=row['jh']
    for r,role in zip(range(74,77),roles): h(f'C{r}',comp[role],'source/chiffrage/chiffrage-sap.yaml')
    for address in ('F72','F78'): h(address,0,'source/plan-investissement.yaml#comparaison_prix_integration_miroir_sap/limite')
    h('C79',p['pilotage_complements_pct']/100,'source/calibration-couts.yaml#pilotage')
    h('F80',p['preparation_eur_ht']/1000,'source/calibration-couts.yaml#couts_communs')
    for r,key in [(84,'bas'),(85,'haut')]: h(f'E{r}',data['scenario']['comparaison_prix_integration_miroir_sap']['contrepoint_integrateur_eur_jh'][key],'source/plan-investissement.yaml#comparaison_prix_integration_miroir_sap')
    for i in range(len(data['options'])):
        for year in range(1,data['scenario']['horizon_ans']+1):
            bind(data['glossaire']['vues']['calculs']['libelle'],f'{chr(66+year)}{6+16*i}',year,'source/plan-investissement.yaml#horizon_ans')
    return out


def check_excel(path,data):
    from openpyxl import load_workbook
    book=load_workbook(path,data_only=False); values=load_workbook(path,data_only=True)
    errors=[]; bound=bindings(data); literal_count=0; tested=0
    norm=lambda x:' '.join(str(x).split())
    known_text=set(); presentation_differences=[]; tooltip_matches=0; formula_matches=0
    for (sheet,address),spec in bound.items():
        actual=values[sheet][address].value; expected=spec['value']
        if expected is None:
            if actual is not None: errors.append(f'{sheet}!{address}: case vide attendue, valeur={actual!r}')
        elif not isinstance(actual,(int,float)) or not math.isclose(actual,expected,rel_tol=1e-12,abs_tol=0.00002 if sheet=='Hypothèses' else 0.00000005):
            errors.append(f'{sheet}!{address}: {actual!r} != YAML {expected!r} ({spec["reference"]})')
        tested+=1
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,(int,float)):
                    literal_count+=1
                    if (sheet.title,cell.coordinate) not in bound: errors.append(f'Constante numérique sans rattachement : {sheet.title}!{cell.coordinate}={cell.value}')
                if cell.data_type=='f' and '[' in cell.value: errors.append(f'Formule externe à examiner : {sheet.title}!{cell.coordinate}')
    result=project(data)
    for i,option in enumerate(result['options']):
        c='CDEFG'[i]
        for year in option['annees']:
            actual=values['Investissements'][f'{c}{6+year["annee"]}'].value
            if not isinstance(actual,(int,float)) or abs(actual-year['total']/1e6)>0.00000005: errors.append(f'Dépense annuelle divergente : {option["stack"]}, année {year["annee"]}')
            tested+=1
    # Resolve the same narrative templates as the generator without exporting Excel.
    cached=json.loads((ROOT/'tmp/investissement-10-ans/inputs.json').read_text(encoding='utf-8'))
    if cached!=data: errors.append('Extraction temporaire du projecteur obsolète ; relancer project_investment.py')
    else:
        proc=subprocess.run([shutil.which('node') or 'node','tools/build_investment_excel.mjs','--check-content','--json'],cwd=ROOT,text=True,encoding='utf-8',capture_output=True)
        if proc.returncode: errors.append('Explications du générateur invalides : '+proc.stderr[-700:])
        else:
            resolved=json.loads(proc.stdout)
            for i,row in enumerate(resolved['perimetre_compare']):
                for c,key in zip('BCDE',('fonction','make','microsoft','sap')):
                    if norm(values['Périmètre'][f'{c}{7+i}'].value)!=norm(row[key]): errors.append(f'Périmètre!{c}{7+i}: texte absent ou divergent du YAML explicatif')
                    known_text.add(('Périmètre',f'{c}{7+i}'))
                    tested+=1
            for i,row in enumerate(resolved['references']):
                for c,key in [('B','libelle'),('C','texte')]:
                    if norm(values['Sources'][f'{c}{28+2*i}'].value)!=norm(row[key]): errors.append(f'Sources!{c}{28+2*i}: référence divergente du YAML explicatif')
                    tested+=1
        proc=subprocess.run([shutil.which('node') or 'node','tools/build_investment_excel.mjs','--check-model','--json'],cwd=ROOT,text=True,encoding='utf-8',capture_output=True)
        if proc.returncode: errors.append('Assemblage en mémoire invalide : '+proc.stderr[-700:])
        else:
            generated=json.loads(proc.stdout)
            captions={(item['sheet'],item['address']):item['value'] for item in generated['captions']}
            known_text.update(captions)
            for (sheet,address),expected in captions.items():
                actual=values[sheet][address].value
                if norm(actual)!=norm(expected):
                    presentation_differences.append(dict(sheet=sheet,cell=address,saved=actual,canonical=expected))
            for tip in generated['tooltips']:
                link=book[tip['sheet']][tip['address']].hyperlink
                if not link or norm(link.tooltip)!=norm(tip['body']):
                    errors.append(f"Infobulle absente ou divergente : {tip['sheet']}!{tip['address']}")
                else: tooltip_matches+=1
            expected_formulas={(item['sheet'],item['address']):item['value'] for item in generated['formulas']}
            saved_formulas={(sheet.title,cell.coordinate):cell.value for sheet in book for row in sheet for cell in row if cell.data_type=='f'}
            def normalize_formula(value):
                # Native Excel removes optional quotes around these local sheet names.
                for view in data['glossaire']['vues'].values():
                    value=value.replace("'"+view['libelle']+"'!",view['libelle']+'!')
                return value
            for key in saved_formulas.keys() | expected_formulas.keys():
                if key not in saved_formulas or key not in expected_formulas or normalize_formula(saved_formulas[key])!=normalize_formula(expected_formulas[key]):
                    errors.append(f'Formule sans rattachement ou divergente : {key[0]}!{key[1]}')
                else: formula_matches+=1
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value,str) and cell.data_type!='f' and (sheet.title,cell.coordinate) not in known_text:
                            errors.append(f'Texte sans rattachement : {sheet.title}!{cell.coordinate}')
    tips=[(s.title,c.coordinate,c.hyperlink.tooltip) for s in book for row in s for c in row if c.hyperlink]
    if len(tips)!=1282 or any(not t or len(t)>255 for _,_,t in tips): errors.append('Couverture ou longueur des infobulles incorrecte')
    book.close();values.close()
    return dict(workbook=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),numeric_literals=literal_count,
                model_bindings=len(bound),checked_cells=tested,text_literals=len(known_text),tooltips=len(tips),tooltip_matches=tooltip_matches,formula_matches=formula_matches,
                presentation_differences=presentation_differences,
                bindings=[dict(sheet=s,cell=a,**v) for (s,a),v in bound.items()]),errors


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',type=Path,default=WORKBOOK)
    parser.add_argument('--json',type=Path,default=ROOT/'tmp/audit-base-connaissance-2026-10-24/controle.json')
    args=parser.parse_args()
    info,errors=inventory(active_files())
    data=prepare();errors+=check_glossary_references(data['glossaire'])
    excel,excel_errors=check_excel(args.workbook,data);errors+=excel_errors
    report=dict(inventory=info,excel=excel,errors=errors,
                limits=['Les références publiques sont contrôlées sur les preuves locales ; aucune nouvelle certification de prix marché.',
                        'Un rattachement technique ne valide pas économiquement les JH, gains de productivité ou transpositions de périmètre.',
                        'Les archives sont conservées à titre historique ; les chiffres calculés ne sont pas des hypothèses autonomes.'])
    args.json.parent.mkdir(parents=True,exist_ok=True)
    args.json.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(dict(files=info['files'],numeric_literals=excel['numeric_literals'],model_bindings=excel['model_bindings'],checked_cells=excel['checked_cells'],text_literals=excel['text_literals'],tooltip_matches=excel['tooltip_matches'],formula_matches=excel['formula_matches'],presentation_differences=len(excel['presentation_differences']),errors=errors),ensure_ascii=False,indent=2))
    if errors: raise SystemExit(1)


if __name__=='__main__': main()
